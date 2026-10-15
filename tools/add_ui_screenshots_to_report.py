from __future__ import annotations

from pathlib import Path
import shutil

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph


SOURCE = Path(r"C:\Users\L2791\Downloads\7.4_PyThinFilm作品说明书_平台更新版.docx")
OUTPUT = Path(r"C:\Users\L2791\Downloads\7.5_PyThinFilm作品说明书_界面截图版.docx")
IMAGE_DIR = Path(
    r"C:\Users\L2791\Downloads\PyThinFilm\Frontend\V3\frontend\report-screenshots"
)

FIGURES = [
    (
        "01_v3_workspace.png",
        "图 4-35：PyThinFilm V3 Workspace 膜系编辑与计算结果总览",
        "PyThinFilm V3 工作区，左侧为膜系编辑和三维结构，右侧为计算状态及光谱指标。",
    ),
    (
        "02_materials_and_3d.png",
        "图 4-36：材料属性折叠面板与可交互三维膜层结构",
        "材料属性面板展开后显示材料类别、真实 n、k 数据和有效波段，三维膜层结构位于下方。",
    ),
    (
        "03_optiland_system.png",
        "图 4-37：Optiland 光学系统二维与三维布局联动界面",
        "Optiland 系统分析界面，展示当前膜系对应的二维光路、三维系统和可折叠的 Spot、PSF、MTF 卡片。",
    ),
    (
        "04_case_library.png",
        "图 4-38：40 例案例库、应用线路与参数输入界面",
        "40 例案例库界面，左侧包含应用线路、系统实验草案和受控参数，右侧为统一输出面板。",
    ),
    (
        "05_rcwa_result.png",
        "图 4-39：RCWA 专用求解器结果与衍射功率曲线",
        "一维周期结构案例的 RCWA 结果界面，显示求解器、关键参数以及 R、T、A 功率曲线。",
    ),
    (
        "06_tamm_field.png",
        "图 4-40：GeneralTmm 真实界面场分布与物理校验",
        "Tamm 界面态的真实归一化场强分布，同时展示共振波长、峰值深度、能量守恒和切向电场连续性检验。",
    ),
]


def find(doc: Document, text: str) -> Paragraph:
    for paragraph in doc.paragraphs:
        if paragraph.text.strip() == text:
            return paragraph
    raise ValueError(f"Paragraph not found: {text}")


def insert_paragraph_before(anchor: Paragraph, style: str = "Normal") -> Paragraph:
    xml = OxmlElement("w:p")
    anchor._p.addprevious(xml)
    paragraph = Paragraph(xml, anchor._parent)
    paragraph.style = style
    return paragraph


def add_picture_before(anchor: Paragraph, image_path: Path, alt_text: str) -> Paragraph:
    paragraph = insert_paragraph_before(anchor)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.keep_with_next = True
    inline_shape = paragraph.add_run().add_picture(str(image_path), width=Inches(5.7))
    inline_shape._inline.docPr.set("descr", alt_text)
    inline_shape._inline.docPr.set("title", image_path.stem)
    return paragraph


def main() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    missing = [name for name, _, _ in FIGURES if not (IMAGE_DIR / name).exists()]
    if missing:
        raise FileNotFoundError(f"Missing screenshots: {missing}")

    shutil.copy2(SOURCE, OUTPUT)
    doc = Document(OUTPUT)
    chapter5 = find(doc, "五、作品分析")

    heading = insert_paragraph_before(chapter5, "Heading 2")
    heading.add_run("4.5 软件界面与计算结果展示")
    heading.paragraph_format.page_break_before = True
    heading.paragraph_format.keep_with_next = True

    intro = insert_paragraph_before(chapter5)
    intro.add_run(
        "本节截图来自当前 PyThinFilm V3 Workspace 的实际运行界面。截图依次展示自由膜系设计、"
        "材料与三维结构、Optiland 光学系统、40 例案例库以及 RCWA、GeneralTmm 专用求解结果，"
        "用于说明前端信息架构、计算来源和物理校验已经形成统一工作流。"
    )
    intro.paragraph_format.keep_with_next = True

    for index, (filename, caption_text, alt_text) in enumerate(FIGURES):
        if index in (2, 4):
            page_break = insert_paragraph_before(chapter5)
            page_break.add_run().add_break(WD_BREAK.PAGE)
        add_picture_before(chapter5, IMAGE_DIR / filename, alt_text)
        caption = insert_paragraph_before(chapter5, "Caption-CN")
        caption.add_run(caption_text)
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.paragraph_format.space_before = Pt(2)
        caption.paragraph_format.space_after = Pt(6)
        caption.paragraph_format.keep_together = True

    doc.core_properties.title = "基于 Python 的薄膜光学仿真设计与教学平台作品说明书"
    doc.core_properties.subject = "PyThinFilm V3 Workspace 平台理论案例验证与软件界面说明"
    doc.core_properties.comments = "在 7.4 平台更新版基础上补充实际软件界面截图。"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
