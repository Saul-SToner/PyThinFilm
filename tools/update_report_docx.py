from __future__ import annotations

from copy import deepcopy
from pathlib import Path
import shutil

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from docx.text.paragraph import Paragraph


SOURCE = Path(r"C:\Users\L2791\Downloads\7.3_第二章理论推导重构版_公式美化严肃版.docx")
OUTPUT = Path(r"C:\Users\L2791\Downloads\7.4_PyThinFilm作品说明书_平台更新版.docx")


def find(doc: Document, text: str) -> Paragraph:
    for paragraph in doc.paragraphs:
        if paragraph.text.strip() == text.strip():
            return paragraph
    raise ValueError(f"Paragraph not found: {text}")


def find_contains(doc: Document, text: str) -> Paragraph:
    for paragraph in doc.paragraphs:
        if text in paragraph.text:
            return paragraph
    raise ValueError(f"Paragraph containing text not found: {text}")


def set_text(paragraph: Paragraph, text: str, style: str | None = None) -> Paragraph:
    paragraph.text = text
    if style:
        paragraph.style = style
    return paragraph


def insert_before(anchor: Paragraph, text: str, style: str = "Normal") -> Paragraph:
    new_xml = OxmlElement("w:p")
    anchor._p.addprevious(new_xml)
    paragraph = Paragraph(new_xml, anchor._parent)
    paragraph.style = style
    paragraph.add_run(text)
    return paragraph


def insert_after(anchor: Paragraph, text: str, style: str = "Normal") -> Paragraph:
    new_xml = OxmlElement("w:p")
    anchor._p.addnext(new_xml)
    paragraph = Paragraph(new_xml, anchor._parent)
    paragraph.style = style
    paragraph.add_run(text)
    return paragraph


def remove(paragraph: Paragraph) -> None:
    parent = paragraph._element.getparent()
    if parent is not None:
        parent.remove(paragraph._element)


def remove_paragraph_border(paragraph: Paragraph) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    border = ppr.find(qn("w:pBdr"))
    if border is not None:
        ppr.remove(border)


def normalize_heading_styles(doc: Document) -> None:
    for name, size in (("Title", 22), ("Heading 1", 16), ("Heading 2", 14), ("Heading 3", 12)):
        style = doc.styles[name]
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.size = Pt(size)
        style.font.bold = True
    title_ppr = doc.styles["Title"].element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)


def compact_range(doc: Document, start_text: str, end_text: str | None, size: float) -> None:
    paragraphs = list(doc.paragraphs)
    start = next(i for i, p in enumerate(paragraphs) if p.text.strip() == start_text)
    end = len(paragraphs)
    if end_text is not None:
        end = next(i for i, p in enumerate(paragraphs[start + 1 :], start + 1) if p.text.strip() == end_text)
    for paragraph in paragraphs[start:end]:
        if paragraph.style.name not in {"Heading 1", "Heading 2"}:
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            for run in paragraph.runs:
                run.font.size = Pt(size)


def add_overview(doc: Document) -> None:
    title = doc.paragraphs[0]
    set_text(title, "基于 Python 的薄膜光学仿真设计与教学平台作品说明书", "Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    remove_paragraph_border(title)
    for run in title.runs:
        run.font.color.rgb = RGBColor(0, 0, 0)

    subtitle = doc.paragraphs[1]
    cursor = subtitle
    blocks = [
        ("作品概述", "Heading 2"),
        (
            "PyThinFilm 面向本科薄膜光学教学与工程入门训练，围绕“理解规律—认识应用—完成创造”的学习路径，将经典膜系理论、真实材料数据、工程案例和光学系统表现组织在同一工作区中。平台当前采用 V3 Workspace 信息架构，通过默认折叠的参数卡片降低信息密度，同时保留膜系结构、主结果图和物理解释的清晰主线。",
            "Normal",
        ),
        (
            "计算层采用按物理问题路由的多求解器体系：平面多层膜和切片渐变结构由 TMMCore 处理；Tamm 界面态的复反射相位与一维真实场由 GeneralTmm 处理；周期光栅由 RCWA 处理；被动辐射制冷光谱由 WPTherml 处理；成像系统布局、Spot、MTF 与 PSF 由 Optiland 负责。外部 COMSOL 数据仅作为严格校验后的正式证据导入，不等同于平台内置有限元求解。",
            "Normal",
        ),
        (
            "当前教学主树共 40 个案例，已形成“可交互主案例、专用求解器案例、正式证据案例”三级交付方式。每个案例均明确输入参数、求解器来源、结果图和模型边界，使学生不仅能够看到曲线，还能够判断曲线由何种模型产生、何时可信以及如何迁移到镜头镀膜等工程任务中。",
            "Normal",
        ),
    ]
    for text, style in blocks:
        cursor = insert_after(cursor, text, style)


def update_toc(doc: Document) -> None:
    anchor = find(doc, "• 2.4 外部仿真数据导入说明")
    insert_after(anchor, "• 2.5 多求解器分层与模型边界", "TocLine-CN")
    anchor = find(doc, "• 2.5 多求解器分层与模型边界")
    insert_after(anchor, "• 2.6 Tamm 界面态与真实场分布", "TocLine-CN")

    replacements = {
        "• 七、版本冻结与后续计划": "• 七、当前版本与后续计划",
        "• 7.1 v0.2.1 竞赛展示功能冻结说明": "• 7.1 当前版本范围",
        "• 7.2 冻结内容": "• 7.2 已完成内容",
        "• 7.3 允许继续修改的内容": "• 7.3 可继续完善内容",
        "• 7.4 不建议加入的新模块": "• 7.4 暂不纳入主线的内容",
        "• 6.5 亚波长 EMT 近似模块的教学扩展价值": "• 6.5 专用求解器与模型边界教学",
    }
    for old, new in replacements.items():
        set_text(find(doc, old), new, "TocLine-CN")


def update_theory_and_architecture(doc: Document) -> None:
    set_text(
        find_contains(doc, "当前竞赛展示版本的核心物理计算仍以一维传输矩阵法"),
        "需要明确的是，外部数据导入与对比绘制仅作为教学对照、计算验证和补充参考。当前版本已在一维 TMM 主线之外接入 RCWA、GeneralTmm 与 WPTherml 等专用求解器，但这些求解器仍各自服从明确的结构维度、材料模型和数值参数边界。外部 COMSOL 数据继续作为独立的参考证据导入，不被表述为平台内部的有限元能力。",
    )

    chapter3 = find(doc, "三、平台开发流程及介绍")
    additions = [
        ("2.5 多求解器分层与模型边界", "Heading 2"),
        (
            "平台不再采用“一个求解器覆盖所有结构”的表述，而是根据结构几何与目标物理量进行求解器路由。TMMCore 用于平面分层介质的 R/T/A 光谱、偏振与角度扫描；Rugate 连续渐变膜、蛾眼渐变结构和 EMT 有效介质在离散切片后仍可调用 TMMCore，但离散步长必须满足收敛性要求。",
            "Normal",
        ),
        (
            "对于横向周期结构，平台调用 RCWA 计算有限傅里叶级次下的衍射与零级光谱；对于 Tamm 金属—DBR 界面，平台调用 GeneralTmm 计算复反射相位和层内前、后向波，从而得到真实的一维场分布；对于 PDRC，WPTherml 用于生成太阳波段与中红外波段的材料光谱响应。Optiland 位于更高层的光学系统表达环节，用于展示镜头布局、Spot、MTF 和 PSF，不替代薄膜求解器。",
            "Normal",
        ),
        (
            "求解器选择遵循三条原则：第一，几何维度与算法假设一致；第二，输出量必须由求解器真实计算而非用示意曲线替代；第三，界面必须同时展示结果来源、关键参数和适用边界。由此避免将切片 TMM、RCWA、有限元证据和成像分析混为同一精度等级。",
            "Normal",
        ),
        ("2.6 Tamm 界面态与真实场分布", "Heading 2"),
        (
            "Tamm 光学态形成于金属与介质布拉格反射镜的分界面。除反射谱中的窄谷外，还应同时考察金属侧与 DBR 侧复反射系数的相位关系。当 φDBR(λ)+φmetal(λ) 接近 2mπ 时，界面附近可形成局域态。平台以反射率极小值和相位匹配共同筛选候选共振波长，避免仅凭曲线形状判断。",
            "Normal",
        ),
        ("IE(z) = |E(z)|² / |Einc|²", "FormulaBlock-CN"),
        (
            "场分布由 GeneralTmm 的层内前向波与后向波系数计算，并按入射场强归一化。界面同时给出共振波长与非共振参考波长的 |E|² 曲线、膜层边界和金属/DBR 分界位置。可靠性检查包括能量守恒误差 |R+T+A−1| 以及界面切向电场连续性误差。",
            "Normal",
        ),
        (
            "在当前保存的 Ag 30 nm + (HL)³H 结构和 400–800 nm 扫描条件下，共振候选位于约 635 nm，非共振参考波长约为 525 nm；归一化峰值 |E|² 约为 3.42，峰值深度约 77.26 nm，金属/DBR 分界位于 30 nm。该数值仅对应当前材料与厚度配置，不作为普适常数。当前算例的能量守恒误差为数值零量级，最大切向电场跳变约 8.5×10⁻⁷。",
            "Normal",
        ),
    ]
    for text, style in additions:
        insert_before(chapter3, text, style)

    set_text(chapter3, "三、平台开发流程与软件架构", "Heading 1")

    chapter4 = find(doc, "四、作品功能与案例展示")
    for paragraph in list(doc.paragraphs):
        if paragraph._p.getparent() is chapter4._p.getparent() and paragraph.text == "" and paragraph._p.sourceline:
            pass
    blocks = [
        (
            "平台开发按“理论核验—接口封装—案例接入—前端呈现—自动验收”的顺序推进。首先以单层减反膜、F-P 滤光片和高反膜作为解析基准；随后将材料、膜系、扫描参数和结果统一为服务接口；最后通过案例目录、前端工作区和测试矩阵验证每一条教学路径。",
            "Normal",
        ),
        (
            "开发过程坚持计算层与展示层分离。求解器只返回可追溯的数值结果和元数据，前端负责曲线、膜层结构、物理解释与交互控件。这样既便于替换或扩展专用求解器，也防止界面为了“有图”而生成缺少物理来源的替代曲线。",
            "Normal",
        ),
        ("3.2 软件目录结构", "Heading 2"),
        ("• thinfilm/：TMMCore、材料、案例目录、接口与验证主模块。", "TocLine-CN"),
        ("• guided_grating/：RCWA、光栅波导及 COMSOL 扫描数据读取支线。", "TocLine-CN"),
        ("• Frontend/V3/backend/：V3 API、专用求解器适配器与案例验收服务。", "TocLine-CN"),
        ("• Frontend/V3/frontend/：V3 Workspace 前端、膜系编辑器、材料检查器和 Optiland 视图。", "TocLine-CN"),
        ("• web3d/public/results/evidence/：正式外部证据、清单与哈希校验信息。", "TocLine-CN"),
        ("• tests/：TMM、材料、案例、RCWA、Tamm、PDRC 和应用层测试。", "TocLine-CN"),
        ("3.3 核心模块说明", "Heading 2"),
        (
            "核心模块由五部分构成：膜系与材料服务负责输入规范化；TMMCore 服务负责平面膜系与切片渐变结构；专用求解器适配器负责 GeneralTmm、RCWA 和 WPTherml；案例目录与验收矩阵负责 40 个案例的能力声明；Optiland 桥接层负责将膜系变化映射到镜头系统的布局与成像指标视图。",
            "Normal",
        ),
        ("3.4 数据流与计算流程", "Heading 2"),
        (
            "一次计算的数据流为：用户选择案例或编辑膜系 → 后端校验材料、厚度、波长、角度与偏振 → 根据结构类型选择求解器 → 生成标准化光谱、场分布或成像结果 → 前端绘图并附加物理解释与模型边界。外部 COMSOL 文件还需经过列名、单位、数值范围和 SHA-256 清单校验后才能进入正式结果区。",
            "Normal",
        ),
        ("3.5 使用方法", "Heading 2"),
        (
            "教学模式下，学生先从案例目录进入任务，展开“实验变量”卡片修改参数，运行后在主图查看曲线或场分布，再展开结果解读、对照指标和模型边界。设计工作区中，用户可编辑膜层顺序、材料和厚度，并通过材料属性下拉项查看 n、k 与数据来源。三维膜层结构支持缩放、拖动和全屏查看。",
            "Normal",
        ),
        (
            "Optiland 页面将二维/三维光学系统布局作为始终可见的主图，Spot、MTF、PSF 和对照信息采用折叠卡片。膜层或材料更新后，系统重新计算关联结果；R/T/A 仍由 PyThinFilm 主工作区承担，避免在光学系统页面重复展示。",
            "Normal",
        ),
        ("3.6 测试与可靠性验证", "Heading 2"),
        (
            "可靠性验证分为解析基准、数值一致性、案例验收和前端构建四层。TMM 核心通过极限结构、能量守恒和向量化一致性测试；专用求解器分别检查相位、场连续性、衍射能量和热辐射光谱；40 个案例均有验收状态；当前 V3 后端专项测试 43 项通过，前端类型检查与生产构建通过。完整仓库测试以 tests/ 目录和 README.md 中的运行结果为准。",
            "Normal",
        ),
    ]
    for text, style in blocks:
        insert_before(chapter4, text, style)

    # Remove the empty placeholders that originally made Chapter 3 a blank page.
    start = find(doc, "3.1 开发流程")
    candidates = list(doc.paragraphs)
    start_idx = next(i for i, paragraph in enumerate(candidates) if paragraph._p is start._p)
    for paragraph in candidates[start_idx + 1 : start_idx + 4]:
        if not paragraph.text.strip():
            remove(paragraph)


def update_cases_and_analysis(doc: Document) -> None:
    set_text(
        find(doc, "本章围绕平台已实现的薄膜光学仿真与前沿应用功能展开。与传统物理理论的静态案例罗列不同，本平台作为可视化教学仿真平台，其核心价值在于支持物理参数的动态交互与可变性调整。本章将仿真案例划分为基础薄膜干涉、真实材料色散、典型工程应用与拓展研究示例四大板块；每一板块均引入原理结构与公式框架，并通过参数扫描展示结构厚度、折射率、周期数、偏振态及物性变化对光谱响应和空间驻波场强的物理调控。为了便于总览，以下首先建立平台案例体系的全局清单表。"),
        "本章围绕教学主树的 40 个案例展开。案例以“物理规律—工程应用—设计任务”为统一叙事，并按完整交互案例、专用求解器案例和正式证据案例三种交付形态接入。所有案例均给出结果图、求解器来源、参数入口与适用边界；其中光谱、场分布、RCWA、PDRC 和成像指标由对应物理模块计算，不使用仅用于占位的替代模型曲线。",
    )
    set_text(
        find(doc, "为了便于全面掌握仿真平台的核心计算资源，以下对平台已实现的多个代表性教学案例、参数扫参案例、真实材料数据库及前沿拓展研究示例进行全局梳理，汇总为“平台案例体系总览表”："),
        "平台案例总览表用于说明 40 个案例分别对应的学习目标、工程对象、求解器与交付状态。当前验收矩阵统计为：18 个完整交互案例、9 个专用求解器案例、13 个正式证据案例，待接入项为 0。该分类反映交互深度与证据来源，不代表案例重要性的高低。",
    )

    replacements = {
        "（2）基准结构与核心公式：结构为：空气 / 40 nm 金（Au）薄膜 / [H/L]^13 DBR / 玻璃基底。在分界面两侧发生局域匹配时，DBR 反射相位角 φ_DBR 与金属反射相位角 φ_metal 应满足共振条件：φ_DBR(λ) + φ_metal(λ) = 2mπ。在带隙内形成强局域界面态，产生极深的光谱吸收谷。": "（2）基准结构与核心公式：当前保存结构为：空气 / 30 nm Ag 薄膜 / (H/L)³H DBR / 基底。GeneralTmm 同时计算两侧复反射系数、R/T/A 和层内场。当 φDBR(λ)+φmetal(λ)=2mπ 的相位条件与反射率极小值在同一波段出现时，将其识别为 Tamm 共振候选。",
        "（3）参数可变性设计：参数可变性表现为改变两侧的边界条件（如调整 DBR 的低折射率层二氧化硅 d_L 物理厚度偏离设计中心），观察界面共振相位匹配波长的位移。": "（3）参数可变性设计：虚拟实验开放 Ag 厚度、DBR 周期数、层厚缩放系数 β、p/s 偏振与扫描波段。学生可观察反射谷、相位匹配与界面场局域峰如何共同移动，并比较共振与非共振场分布。",
        "图 4-29：不同 DBR 厚度偏离下 Tamm 界面态局域共振吸收对比图（参数扫描图位）": "图 4-29：Tamm 反射谱、相位匹配与参数扫描结果",
        "（5）结果与讨论：由图 4-30 可以看出，随着 DBR 厚度发生 5% 的微小偏离，平台计算得到的反射带内部局域共振吸收峰发生位移，吸收峰形也随之改变。该现象与 Tamm 界面局域相位匹配零和物理公式相一致，说明局域共振波长主要由两侧复反射系数的相位延迟决定。比对不同 DBR 参数下的吸收谷位移曲线，能够让学生清晰体会到电磁局域态对反射相位边界条件的较高敏感性，将电磁波动相位转换为直观的频谱透射特征。": "（5）结果与讨论：当前默认结构在 400–800 nm 扫描中得到约 635 nm 的共振候选和约 525 nm 的非共振参考点。共振场的归一化峰值 |E|² 约为 3.42，峰值位于约 77.26 nm 深度，且在 30 nm 的 Ag/DBR 分界附近呈现明显局域。能量守恒与切向电场连续性检查同步给出，使学生能够把反射谷、相位条件和空间场增强对应起来。",
        "图 4-30：Tamm 共振局域极化界面态吸收与反射相位角解算图": "图 4-30：Tamm 共振与非共振状态下的真实一维 |E|² 场分布",
        "本项目开发的平台不是商业仿真软件（如有限元软件 COMSOL 等）的替代品，而是作为教学演示与初步设计分析的轻量化补充方案。对于平面多层膜及亚波长周期结构的零阶估算，本平台具有计算时间短、易部署的优势；但对于大倾角共振、任意二维/三维周期结构，仍需配合主流商业软件进行全波物理计算。": "本平台不是 COMSOL、Zemax 等商业软件的替代品，而是面向教学解释、快速验证与初步设计的开放式补充。TMMCore 适合平面膜系，RCWA 适合一维周期结构，GeneralTmm 适合平面复相位与场分布，WPTherml 适合热辐射光谱，Optiland 负责光学系统级成像分析。对于任意二维/三维全波结构、粗糙散射和完整制造公差分析，仍需使用相应的专业工具。",
        "有效介质近似模块假设周期远小于波长（P ≪ λ），仅计算零阶极化干涉。若周期参数与工作波长接近，由于发生波导共振、强多阶衍射激发，EMT 模型不再适用，需要采用更高阶的周期结构电磁求解方法。": "EMT 仅在结构尺度远小于波长时作为零阶近似；周期接近波长后应切换到 RCWA。当前 RCWA 实现面向一维周期结构，精度受傅里叶级次、材料色散和收敛设置影响，不能据此宣称覆盖任意二维/三维光栅。",
        "平台外部数据读取管线仅用于导入外部仿真生成的 CSV 离散数据并进行特征提取，平台内部计算引擎并不支持全波方程的数值离散物理求解。": "平台外部数据读取管线仅用于导入 COMSOL 等软件生成的 CSV 离散数据并进行特征提取。正式证据文件需通过模板、单位、数值范围和 SHA-256 校验。该功能不代表平台内部拥有有限元求解器，也不应与内置 RCWA 结果混淆。",
    }
    for old, new in replacements.items():
        set_text(find(doc, old), new)


def update_innovation_roadmap_and_appendix(doc: Document) -> None:
    replacements = {
        "平台打破了传统教学软件仅有演示的局限，实现了“基础多层膜物理现象解析 → 工业工程案例分析 → 外部参考曲线读取验证”的一维光学薄膜闭环分析设计流程。": "平台以“理解规律 → 认识应用 → 完成创造”为学习主线，将物理案例、镜头镀膜等工程应用和开放设计任务连接起来。学生既能修改参数观察结果，也能追踪求解器、材料数据和模型边界。",
        "采用自动化测试套件（270 个 pytest 用例），对多层膜极限退化、能量守恒定律进行断言检验，提高了计算流程的可复现性与数值稳定性。": "采用分层自动化测试覆盖 TMM 极限退化、向量化一致性、真实材料、案例路由、RCWA、Tamm、PDRC 与工程应用。当前 V3 后端专项测试 43 项通过，40 个教学案例均进入验收矩阵；完整测试数量以仓库 tests/ 目录的实时结果为准。",
        "6.5 亚波长 EMT 近似模块的教学扩展价值": "6.5 专用求解器与模型边界教学",
        "利用 EMT 零阶极化等效公式简化微纳格栅计算，使本科生可在不掌握复杂波级耦合矩阵的前提下，提前接触偏振各向异性、双折射干涉等微纳结构物理概念。": "在 EMT 直观近似之外接入 RCWA、GeneralTmm 与 WPTherml，使学生能够比较不同模型的假设、输入和输出。界面只展示真实计算结果，并明确说明收敛条件与不能覆盖的结构类型。",
        "提供 7 种 Plotly 交互图表，能以单包 html 导出，为多媒体教学提供交互与展示支持。": "V3 Workspace 采用低信息密度布局：膜系结构与主结果常显，材料属性、实验变量、Spot、MTF、PSF 和对照指标按需折叠。膜层三维视图支持缩放、拖动和全屏，兼顾课堂演示与自主实验。",
        "七、版本冻结与后续计划": "七、当前版本与后续计划",
        "7.1 v0.2.1 竞赛展示功能冻结说明": "7.1 当前版本范围",
        "为保证在竞赛演示和评审过程中的稳定性，本项目版本确立为 v0.2.1。": "当前主线由教学仿真主树与光栅波导研究支线构成。教学平台不开放厚度反演入口，仓库不保存反演样本 CSV；当前文档以 V3 Workspace 和 40 个案例验收矩阵为功能基线。",
        "本版本正式进入功能冻结保护，以防止大范围逻辑重构带来不可预测的计算偏离或异常。": "版本迭代遵循“计算证据先于界面声明”的原则：只有求解器、测试和结果来源明确的能力才进入主界面。",
        "7.2 冻结内容": "7.2 已完成内容",
        "冻结了新增的物理模块、新的数学物理计算公式、新的计算大类主线，以及除 TMM、EMT、COMSOL CSV 以外的额外电磁求解器开发。": "已完成 V3 工作区、膜系与材料编辑、40 个案例验收、多求解器路由、Tamm 真实场、RCWA、PDRC、Optiland 光学系统视图以及 COMSOL 正式证据导入。",
        "7.3 允许继续修改的内容": "7.3 可继续完善内容",
        "允许对现有文档、相对路径链接进行修正与拼写校正；允许对终端中文编码等 BUG 进行兼容性修复；允许增加新单元测试用例及补充冒烟测试数据；允许优化 Plotly 图表展示样式及课堂演示材料。": "后续可继续完善 RCWA 与 WPTherml 的实验变量、真实色散材料覆盖、求解器收敛提示、工作区性能和课堂任务说明，并保持案例验收与自动化测试同步更新。",
        "7.4 不建议加入的新模块": "7.4 暂不纳入主线的内容",
        "1. 严格耦合波级次矩阵特征值求解器：其傅里叶空间阶数多、矩阵复数本征解在极化方向易发散，容易引起计算不稳定。": "1. 任意二维/三维全波求解：当前 RCWA 只承担一维周期结构计算，不将其扩展表述为通用全波平台。",
        "2. 多层 Web 网页交互式控制台接口：当前版本主要通过 Python 脚本、案例入口和 Plotly 交互图表完成仿真演示。后续可在不改变核心计算模块的基础上封装为轻量化演示界面。不建议在冻结期继续增加网页开发框架依赖包。": "2. Zemax 导入导出与全面 Three.js 迁移：短期继续保留现有原生交互方案，优先稳定计算、信息架构和教学路径。",
        "3. 厚度多参全局逆向设计反演算法：多维反演具有不收敛与解的多值共存性，本版本已彻底移除相关逆向反演，禁止重新加载。": "3. 厚度多参反演入口：研究代码与教学主树保持分离，教学平台不暴露反演功能，也不在仓库中保留反演样本 CSV。",
        "后续计划在核心 TMM 计算引擎保持稳定的前提下，可探索轻量化 Web 展示接口，以提高参数调节的可交互体验，并进一步研究将高阶傅里叶衍射解算器引入 1D 周期微纳结构的物理仿真中。": "后续重点是完善每个案例的“规律—应用—设计”任务书、增加专用求解器的可解释参数扫描、加强材料数据溯源与收敛性提示，并继续用 Optiland 展示膜系变化对真实光学系统的影响。",
    }
    for old, new in replacements.items():
        set_text(find(doc, old), new)

    dep_anchor = find(doc, "• plotly>=5.0")
    for text in [
        "• FastAPI / Uvicorn：V3 后端接口",
        "• React / TypeScript：V3 前端工作区",
        "• tmmcore 0.4.0：平面多层膜计算",
        "• rcwa 1.0.48：一维周期结构计算",
        "• GeneralTmm 1.3.1：Tamm 相位与真实场计算",
        "• wptherml 1.1.1：热辐射与 PDRC 光谱计算",
        "• Optiland：独立环境中的光学系统布局与成像分析",
    ]:
        dep_anchor = insert_after(dep_anchor, text, "TocLine-CN")

    set_text(find(doc, "• thinfilm/education.py：TMM 及拓展物理模型。"), "• thinfilm/：TMMCore、材料、教学案例、API 与验证主模块。", "TocLine-CN")
    set_text(find(doc, "• guided_grating/comsol_io.py：COMSOL CSV 导入器。"), "• guided_grating/：RCWA、光栅波导与 COMSOL CSV 严格读取。", "TocLine-CN")
    appendix_anchor = find(doc, "• examples/applications/：工程案例集合。")
    appendix_anchor = insert_after(appendix_anchor, "• Frontend/V3/backend/：多求解器 API 与案例验收。", "TocLine-CN")
    insert_after(appendix_anchor, "• Frontend/V3/frontend/：V3 Workspace 用户界面。", "TocLine-CN")

    set_text(find(doc, "• rcwa_1d：亚波长周期光栅有效介质理论 (EMT) 零阶等效偏振反射率计算。"), "• rcwa_1d：一维周期光栅 RCWA 求解与级次能量输出。", "TocLine-CN")
    core_anchor = find(doc, "• rcwa_1d：一维周期光栅 RCWA 求解与级次能量输出。")
    core_anchor = insert_after(core_anchor, "• multilayer_field_comparison：共振与非共振真实一维场分布对比。", "TocLine-CN")
    insert_after(core_anchor, "• /api/case-acceptance：40 个案例的交付形态与验收状态接口。", "TocLine-CN")

    run_anchor = find(doc, "• 运行工程案例：python run_case.py --group teaching --case metrics_bundle")
    insert_after(run_anchor, "• 启动 V3 前端：进入 Frontend/V3/frontend 后运行 npm run dev", "TocLine-CN")


def add_footer_metadata(doc: Document) -> None:
    for section in doc.sections:
        footer = section.footer
        if footer.paragraphs:
            p = footer.paragraphs[0]
        else:
            p = footer.add_paragraph()
        if not p.text.strip():
            p.text = "PyThinFilm V3 Workspace｜作品说明书更新版"
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.font.size = Pt(8)
                run.font.color.rgb = RGBColor(128, 128, 128)


def main() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    shutil.copy2(SOURCE, OUTPUT)
    doc = Document(OUTPUT)
    normalize_heading_styles(doc)
    add_overview(doc)
    update_toc(doc)
    update_theory_and_architecture(doc)
    update_cases_and_analysis(doc)
    update_innovation_roadmap_and_appendix(doc)
    compact_range(doc, "八、补充信息说明", "九、参考文献", 9.0)
    compact_range(doc, "十、附录", None, 8.5)
    add_footer_metadata(doc)
    doc.core_properties.title = "基于 Python 的薄膜光学仿真设计与教学平台作品说明书"
    doc.core_properties.subject = "PyThinFilm V3 Workspace 平台、理论、案例与验证说明"
    doc.core_properties.comments = "在 7.3 理论推导重构版基础上更新，保留原稿。"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
