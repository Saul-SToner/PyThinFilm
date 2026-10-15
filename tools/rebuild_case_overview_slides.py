from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(r"C:\Users\L2791\Downloads\PyThinFilm")
INPUT = Path(r"C:\Users\L2791\Downloads\pythinfilm_competition_deck_20260711_页码整理版.pptx")
OUTPUT = Path(r"C:\Users\L2791\Downloads\pythinfilm_competition_deck_20260711_页码整理版_案例原图高清版.pptx")
IMG_DIR = ROOT / "projects" / "pythinfilm_competition_deck_ppt169_20260710" / "images"


SLIDES = [
    {
        "number": 23,
        "labels": [
            "四分之一波长单层减反膜", "半波长单层相位膜", "四分之一波长双层减反膜",
            "多孔SiO₂减反膜", "多孔双层减反膜", "蛾眼减反：结构示意",
            "蛾眼减反：参数扫描", "四分之一波长周期膜", "Bragg高反膜", "DBR-only对照结构",
        ],
    },
    {
        "number": 24,
        "labels": [
            "F-P滤光片", "窄带滤光片：单腔", "窄带滤光片：多层腔", "Rugate渐变折射率滤光片",
            "中性分束薄膜", "TPP反射型近完美吸收结构", "TPP近完美吸收条件",
            "TPP结构电磁场分布", "TPP钨膜厚度扫描", "平面吸收层与粗糙吸收表面对照",
        ],
    },
    {
        "number": 25,
        "labels": [
            "准随机粗糙吸收表面增强吸收", "粗糙表面电磁场热点", "光栅波导基础结构",
            "光栅波导相位匹配机理", "光栅波导谱线与近似对比", "PDRC参数筛选流程",
            "PDRC双波段光谱调控机理", "PDRC被动日间辐射冷却多层膜", "拓扑Tamm边界态阶梯结构",
            "Tamm/TPP反射相位与复平面卷绕",
        ],
    },
]


IMAGE_FILES = [
    "PyThinFilm_作品说明书_p24_11.png", "PyThinFilm_作品说明书_p25_12.png",
    "PyThinFilm_作品说明书_p26_13.png", "PyThinFilm_作品说明书_p27_14.png",
    "PyThinFilm_作品说明书_p28_15.png", "PyThinFilm_作品说明书_p29_16.png",
    "PyThinFilm_作品说明书_p30_17.png", "PyThinFilm_作品说明书_p31_18.png",
    "PyThinFilm_作品说明书_p32_19.png", "PyThinFilm_作品说明书_p32_20.png",
    "PyThinFilm_作品说明书_p34_21.png", "PyThinFilm_作品说明书_p35_22.png",
    "PyThinFilm_作品说明书_p36_23.png", "PyThinFilm_作品说明书_p37_24.png",
    "PyThinFilm_作品说明书_p38_25.png", "PyThinFilm_作品说明书_p39_26.png",
    "PyThinFilm_作品说明书_p41_27.png", "PyThinFilm_作品说明书_p42_28.png",
    "PyThinFilm_作品说明书_p43_29.png", "PyThinFilm_作品说明书_p44_30.png",
    "PyThinFilm_作品说明书_p45_31.png", "PyThinFilm_作品说明书_p46_32.png",
    "PyThinFilm_作品说明书_p47_33.png", "PyThinFilm_作品说明书_p48_34.png",
    "PyThinFilm_作品说明书_p49_35.png", "PyThinFilm_作品说明书_p50_36.png",
    "PyThinFilm_作品说明书_p51_37.png", "PyThinFilm_作品说明书_p52_38.png",
    "PyThinFilm_作品说明书_p53_39.png", "PyThinFilm_作品说明书_p53_40.png",
]


def remove_large_screenshot(slide):
    for shape in list(slide.shapes):
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE and shape.width > Inches(8):
            shape._element.getparent().remove(shape._element)


def add_contained_picture(slide, path, x, y, w, h):
    with Image.open(path) as im:
        iw, ih = im.size
    ratio = min(w / iw, h / ih)
    pw, ph = int(iw * ratio), int(ih * ratio)
    return slide.shapes.add_picture(str(path), x + (w - pw) // 2, y + (h - ph) // 2, pw, ph)


def add_card(slide, idx, image_path, label):
    col, row = idx % 5, idx // 5
    x = Inches(0.48 + col * 2.57)
    y = Inches(1.38 + row * 2.58)
    w, h = Inches(2.42), Inches(2.38)

    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    card.fill.solid(); card.fill.fore_color.rgb = RGBColor(248, 251, 255)
    card.line.color.rgb = RGBColor(147, 187, 229); card.line.width = Pt(1.2)

    add_contained_picture(slide, image_path, x + Inches(0.10), y + Inches(0.18), w - Inches(0.20), Inches(1.62))

    badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x + Inches(0.02), y + Inches(0.02), Inches(0.32), Inches(0.25))
    badge.fill.solid(); badge.fill.fore_color.rgb = RGBColor(49, 132, 216)
    badge.line.fill.background()
    tf = badge.text_frame; tf.clear(); tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = str(idx + 1); r.font.name = "Arial"; r.font.size = Pt(9); r.font.bold = True; r.font.color.rgb = RGBColor(255, 255, 255)

    box = slide.shapes.add_textbox(x + Inches(0.10), y + Inches(1.86), w - Inches(0.20), Inches(0.40))
    tf = box.text_frame; tf.clear(); tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = label; r.font.name = "Microsoft YaHei"; r.font.size = Pt(9.5); r.font.bold = True; r.font.color.rgb = RGBColor(30, 70, 115)


def main():
    missing = [name for name in IMAGE_FILES if not (IMG_DIR / name).is_file()]
    if missing:
        raise FileNotFoundError(f"Missing source images: {missing}")
    prs = Presentation(str(INPUT))
    for group_idx, spec in enumerate(SLIDES):
        slide = prs.slides[spec["number"] - 1]
        remove_large_screenshot(slide)
        start = group_idx * 10
        for idx, label in enumerate(spec["labels"]):
            add_card(slide, idx, IMG_DIR / IMAGE_FILES[start + idx], label)
    prs.save(str(OUTPUT))
    print(OUTPUT)


if __name__ == "__main__":
    main()
